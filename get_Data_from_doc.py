# 从word文件中获取图片以及描述数据，由于docx不支持doc先将其转换
from win32com import client as wc
import docx
import os
import zipfile

filePath = r'C:\Users\16254\Desktop\机器学习\2012年度岩屑薄片鉴定\鉴定文档\\'
fileName = r'C:\Users\16254\Desktop\机器学习\2012年度岩屑薄片鉴定\鉴定文档\D28-1-870.doc'
COUNTS = 0
imageID =[]
def doc2docx(fileName):
    # 首先将doc转换成docx
    word = wc.Dispatch("Word.Application")
    doc = word.Documents.Open(fileName)
    # 使用参数16表示将doc转换成docx，保存成docx后才能 读文件
    FileNameDocx = fileName[:-4] + '.docx'
    doc.SaveAs(FileNameDocx, 16)
    doc.Close()
    word.Quit()
    return FileNameDocx

def get_Class_Docx(FileNameDocx):
    # 获取文档对象docx，提取其中的文字描述
    file = docx.Document(FileNameDocx)
    identify = []
    for table in file.tables:
        for i in range(len(table.rows)):
            for j in range(len(table.columns)):
                identify.append(table.cell(i, j).text.strip().replace('，','+'))

    identify = identify[identify.index('岩屑镜下照片'):identify.index('岩屑镜下描述')]
    imageClass = [x for x in identify if (x != '') and (x != '岩屑镜下照片')]
    return imageClass

def extract_Images(FileNameDocx):
    # 提取docx文件中的图片
    global COUNTS
    FileNameZip = FileNameDocx[:-5] + '.ZIP'
    os.rename(FileNameDocx, FileNameZip)  # 重命名为zip文件
    # 进行解压
    with zipfile.ZipFile(FileNameZip, 'r') as f :
        fileImage = [x for x in f.namelist() if 'word/media/' in x]
        #创建临时目录保存提取内容
        fileDir = filePath+'tmp'
        os.mkdir(fileDir)
        global imageID
        imageID = []
        for file in fileImage:
            f.extract(file,fileDir)
            imageID.append(COUNTS)
            os.rename(fileDir+'\\'+file, filePath+'images\\'+str(COUNTS)+os.path.splitext(file)[1])
            COUNTS += 1
    # 只为删除tmp文件夹保证结果准确
        os.renames(fileDir+'/word/media', FileNameZip[:-4])
        os.removedirs(FileNameZip[:-4])
    os.remove(FileNameZip)
    return len(fileImage)

def annotations_out():
    # 输出描述文件
    with open(filePath + '\\images\\annotation.txt', 'a') as annotations:
        for i,data in enumerate(imageClass):
            rowtxt = '{}:{}\n'.format(imageID[i],data)
            annotations.write(rowtxt)


if __name__ == '__main__':
    #根据路径遍历需要处理的文件
    for name in os.listdir(filePath):
        if 'doc' in name:
            fileName = filePath + name
            FileNameDocx = doc2docx(fileName)
            imageClass = get_Class_Docx(FileNameDocx)
            print('1',imageClass)
            extractNum = extract_Images(FileNameDocx)
            if extractNum == len(imageClass):
                annotations_out()
            else:
                print('描述文件无效')
                continue
        else:
            continue





# shutil.move(fileDir+'/word/media',fileDir)
# os.renames(fileDir+'/word/media', FileNameZip[:-4])
# os.remove(FileNameZip)
